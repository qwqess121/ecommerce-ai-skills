#!/usr/bin/env python3
"""
Listing 诊断分析 Excel 报告生成器

使用方式：
  python generate-report.py --data diagnosis_data.json --output "Listing诊断_2026-09-26.xlsx"

输入：由 Claude skill 生成的 JSON 诊断数据
输出：格式化的 Excel 报告（多 Sheet）
"""

import json
import sys
import os
from datetime import datetime
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import (
        Font, PatternFill, Alignment, Border, Side, numbers
    )
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl not installed. Run: python -m pip install openpyxl")
    sys.exit(1)


AMAZON_HEADERS = [
    ("国家", 8),
    ("平台", 10),
    ("店铺", 15),
    ("负责人", 10),
    ("ASIN", 14),
    ("原标题", 45),
    ("修改标题", 45),
    ("五点描述建议", 60),
    ("主副图修改建议", 40),
    ("A+详情修改建议", 40),
    ("错词", 15),
    ("正确词", 15),
    ("错词位置", 10),
    ("违规词", 15),
    ("违规类别", 12),
    ("严重程度", 10),
    ("建议替换词", 15),
    ("综合评分", 10),
    ("优先修改项", 35),
]

TIKTOK_HEADERS = [
    ("国家", 8),
    ("平台", 12),
    ("店铺", 15),
    ("负责人", 10),
    ("产品ID", 18),
    ("原标题", 45),
    ("修改标题", 45),
    ("详情描述建议", 60),
    ("主副图修改建议", 40),
    ("错词", 15),
    ("正确词", 15),
    ("错词位置", 10),
    ("违规词", 15),
    ("违规类别", 12),
    ("严重程度", 10),
    ("建议替换词", 15),
    ("综合评分", 10),
    ("优先修改项", 35),
]

COMPETITOR_APPEARANCE_HEADERS = [
    ("平台", 10),
    ("我的ASIN/产品ID", 18),
    ("我的标题", 40),
    ("维度", 12),
    ("竞品1", 35),
    ("竞品2", 35),
    ("竞品3", 35),
    ("竞品4", 35),
    ("竞品5", 35),
]

COMPETITOR_CATEGORY_HEADERS = [
    ("平台", 10),
    ("品类/场景", 20),
    ("统计维度", 15),
    ("我的数据", 20),
    ("品类前20均值", 20),
    ("品类前20中位数", 20),
    ("品类前20最小值", 20),
    ("品类前20最大值", 20),
    ("差距分析", 35),
]

VIOLATION_SUMMARY_HEADERS = [
    ("平台", 10),
    ("ASIN/产品ID", 18),
    ("违规词", 15),
    ("违规类别", 15),
    ("严重程度", 10),
    ("出现位置", 15),
    ("建议替换词", 20),
    ("平台规则说明", 40),
]

TYPO_SUMMARY_HEADERS = [
    ("平台", 10),
    ("ASIN/产品ID", 18),
    ("错词", 15),
    ("正确词", 15),
    ("出现位置", 15),
    ("出现次数", 10),
]

AD_STRUCTURE_HEADERS = [
    ("ASIN", 14),
    ("国家", 8),
    ("店铺", 15),
    ("Campaign总数", 12),
    ("SP精准组数", 12),
    ("SP广泛组数", 12),
    ("SP自动组数", 12),
    ("SB Campaign数", 12),
    ("SBV数", 10),
    ("SD数", 10),
    ("月度花费", 12),
    ("平均ACOS", 10),
    ("关键词覆盖数", 12),
    ("出单词数", 10),
    ("vs Top5平均差距", 15),
]

AD_COMPETITOR_HEADERS = [
    ("排名", 8),
    ("竞品ASIN", 14),
    ("Campaign数", 12),
    ("SP精准", 12),
    ("SP广泛", 12),
    ("SP自动", 12),
    ("SB数", 10),
    ("平均ACOS", 10),
    ("关键词覆盖数", 12),
    ("出单词数", 10),
    ("花费Top3Campaign", 35),
    ("vs本品ACOS差距", 15),
]

KEYWORD_OPPORTUNITY_HEADERS = [
    ("竞品ASIN", 14),
    ("高效出单词", 20),
    ("搜索量", 10),
    ("竞品ACOS", 10),
    ("本品是否覆盖", 15),
    ("本品排名", 12),
    ("机会评分", 10),
    ("建议优先级", 12),
]

TRAFFIC_STRUCTURE_HEADERS = [
    ("ASIN", 14),
    ("自然流量%", 12),
    ("广告流量%", 12),
    ("直接流量%", 12),
    ("搜索流量%", 12),
    ("品牌词占比", 12),
    ("类目词占比", 12),
    ("竞品词占比", 12),
    ("流量趋势", 15),
]

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
DATA_FONT = Font(name="Microsoft YaHei", size=10)
SCORE_HIGH_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
SCORE_MED_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
SCORE_LOW_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
SEVERITY_HIGH_FONT = Font(name="Microsoft YaHei", size=10, color="CC0000", bold=True)
SEVERITY_MED_FONT = Font(name="Microsoft YaHei", size=10, color="CC6600")
SEVERITY_LOW_FONT = Font(name="Microsoft YaHei", size=10, color="336600")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
WRAP_ALIGNMENT = Alignment(vertical="top", wrap_text=True)


def setup_sheet(ws, headers):
    for col_idx, (header_name, width) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header_name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def write_data_row(ws, row_idx, data, headers):
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.font = DATA_FONT
        cell.border = THIN_BORDER
        cell.alignment = WRAP_ALIGNMENT

        header_name = headers[col_idx - 1][0]
        value = data.get(header_name, "")
        cell.value = value

        if header_name == "综合评分" and isinstance(value, (int, float)):
            if value >= 80:
                cell.fill = SCORE_HIGH_FILL
            elif value >= 60:
                cell.fill = SCORE_MED_FILL
            else:
                cell.fill = SCORE_LOW_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")

        if header_name == "严重程度":
            if value == "高":
                cell.font = SEVERITY_HIGH_FONT
            elif value == "中":
                cell.font = SEVERITY_MED_FONT
            elif value == "低":
                cell.font = SEVERITY_LOW_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")


def add_summary_sheet(wb, data):
    ws = wb.create_sheet("诊断摘要", 0)

    title_font = Font(name="Microsoft YaHei", size=16, bold=True, color="1F4E79")
    subtitle_font = Font(name="Microsoft YaHei", size=12, bold=True)
    stat_font = Font(name="Microsoft YaHei", size=11)

    ws.merge_cells("A1:F1")
    cell = ws["A1"]
    cell.value = f"Listing 诊断周报 — {data.get('report_date', datetime.now().strftime('%Y-%m-%d'))}"
    cell.font = title_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    row = 3
    for platform_name, platform_key in [("亚马逊", "amazon"), ("TikTok Shop", "tiktok")]:
        stats = data.get(f"{platform_key}_stats", {})
        if not stats:
            continue

        ws.cell(row=row, column=1, value=f"{platform_name} 诊断概览").font = subtitle_font
        row += 1

        stat_items = [
            ("诊断产品数", stats.get("total", 0)),
            ("优秀 (80+)", stats.get("excellent", 0)),
            ("需改进 (60-79)", stats.get("needs_improvement", 0)),
            ("紧急 (<60)", stats.get("urgent", 0)),
            ("违规词检出", f"{stats.get('violations_total', 0)}处（高危{stats.get('violations_high', 0)}处）"),
            ("错词检出", f"{stats.get('typos_total', 0)}处"),
        ]

        for label, value in stat_items:
            ws.cell(row=row, column=1, value=label).font = stat_font
            ws.cell(row=row, column=2, value=str(value)).font = stat_font
            row += 1

        row += 1

    # 广告诊断汇总（仅亚马逊）
    ad_stats = data.get("ad_stats", {})
    if ad_stats:
        ws.cell(row=row, column=1, value="亚马逊 广告诊断概览").font = subtitle_font
        row += 1

        ad_items = [
            ("平均ACOS", f"{ad_stats.get('avg_acos', 'N/A')}%"),
            ("Campaign总数", ad_stats.get("total_campaigns", 0)),
            ("机会词数", ad_stats.get("opportunity_keywords", 0)),
            ("vs Top5竞品ACOS差距", f"{ad_stats.get('vs_top5_gap', 'N/A')}%"),
        ]

        for label, value in ad_items:
            ws.cell(row=row, column=1, value=label).font = stat_font
            ws.cell(row=row, column=2, value=str(value)).font = stat_font
            row += 1

        row += 1

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 30


def generate_report(data, output_path):
    wb = Workbook()
    wb.remove(wb.active)

    add_summary_sheet(wb, data)

    ws_amazon = wb.create_sheet("亚马逊诊断报告")
    setup_sheet(ws_amazon, AMAZON_HEADERS)
    for idx, item in enumerate(data.get("amazon_diagnosis", []), 2):
        write_data_row(ws_amazon, idx, item, AMAZON_HEADERS)

    ws_tiktok = wb.create_sheet("TikTok诊断报告")
    setup_sheet(ws_tiktok, TIKTOK_HEADERS)
    for idx, item in enumerate(data.get("tiktok_diagnosis", []), 2):
        write_data_row(ws_tiktok, idx, item, TIKTOK_HEADERS)

    ws_comp_app = wb.create_sheet("竞品对标-同外观Top5")
    setup_sheet(ws_comp_app, COMPETITOR_APPEARANCE_HEADERS)
    for idx, item in enumerate(data.get("competitor_appearance", []), 2):
        write_data_row(ws_comp_app, idx, item, COMPETITOR_APPEARANCE_HEADERS)

    ws_comp_cat = wb.create_sheet("竞品对标-同场景品类前20")
    setup_sheet(ws_comp_cat, COMPETITOR_CATEGORY_HEADERS)
    for idx, item in enumerate(data.get("competitor_category", []), 2):
        write_data_row(ws_comp_cat, idx, item, COMPETITOR_CATEGORY_HEADERS)

    ws_ad_structure = wb.create_sheet("广告结构拆分（本品）")
    setup_sheet(ws_ad_structure, AD_STRUCTURE_HEADERS)
    for idx, item in enumerate(data.get("ad_structure", []), 2):
        write_data_row(ws_ad_structure, idx, item, AD_STRUCTURE_HEADERS)

    ws_ad_competitor = wb.create_sheet("广告对标矩阵（Top20竞品）")
    setup_sheet(ws_ad_competitor, AD_COMPETITOR_HEADERS)
    for idx, item in enumerate(data.get("ad_competitor", []), 2):
        write_data_row(ws_ad_competitor, idx, item, AD_COMPETITOR_HEADERS)

    ws_keywords = wb.create_sheet("关键词机会词（出单词反查）")
    setup_sheet(ws_keywords, KEYWORD_OPPORTUNITY_HEADERS)
    for idx, item in enumerate(data.get("keyword_opportunities", []), 2):
        write_data_row(ws_keywords, idx, item, KEYWORD_OPPORTUNITY_HEADERS)

    ws_traffic = wb.create_sheet("流量结构对比（本品vs Top5）")
    setup_sheet(ws_traffic, TRAFFIC_STRUCTURE_HEADERS)
    for idx, item in enumerate(data.get("traffic_structure", []), 2):
        write_data_row(ws_traffic, idx, item, TRAFFIC_STRUCTURE_HEADERS)

    ws_violations = wb.create_sheet("违规词汇总")
    setup_sheet(ws_violations, VIOLATION_SUMMARY_HEADERS)
    for idx, item in enumerate(data.get("violations_summary", []), 2):
        write_data_row(ws_violations, idx, item, VIOLATION_SUMMARY_HEADERS)

    ws_typos = wb.create_sheet("错词汇总")
    setup_sheet(ws_typos, TYPO_SUMMARY_HEADERS)
    for idx, item in enumerate(data.get("typos_summary", []), 2):
        write_data_row(ws_typos, idx, item, TYPO_SUMMARY_HEADERS)

    output_dir = os.path.dirname(output_path)
    if output_dir and not os.path.exists(output_dir):
        os.makedirs(output_dir, exist_ok=True)

    wb.save(output_path)
    print(f"Report saved to: {output_path}")
    return output_path


def generate_empty_template(output_path):
    """Generate an empty template Excel with all sheets and headers."""
    data = {
        "report_date": datetime.now().strftime("%Y-%m-%d"),
        "amazon_stats": {"total": 0, "excellent": 0, "needs_improvement": 0, "urgent": 0, "violations_total": 0, "violations_high": 0, "typos_total": 0},
        "tiktok_stats": {"total": 0, "excellent": 0, "needs_improvement": 0, "urgent": 0, "violations_total": 0, "violations_high": 0, "typos_total": 0},
        "ad_stats": {"avg_acos": "N/A", "total_campaigns": 0, "opportunity_keywords": 0, "vs_top5_gap": "N/A"},
        "amazon_diagnosis": [],
        "tiktok_diagnosis": [],
        "ad_structure": [],
        "ad_competitor": [],
        "keyword_opportunities": [],
        "traffic_structure": [],
        "competitor_appearance": [],
        "competitor_category": [],
        "violations_summary": [],
        "typos_summary": [],
    }
    return generate_report(data, output_path)


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Listing Diagnosis Excel Report Generator")
    parser.add_argument("--data", help="Path to diagnosis JSON data file")
    parser.add_argument("--output", help="Output Excel file path", default=f"Listing诊断_{datetime.now().strftime('%Y-%m-%d')}.xlsx")
    parser.add_argument("--template", action="store_true", help="Generate empty template")

    args = parser.parse_args()

    if args.template:
        generate_empty_template(args.output)
        print(f"Empty template generated: {args.output}")
    elif args.data:
        with open(args.data, "r", encoding="utf-8") as f:
            diagnosis_data = json.load(f)
        generate_report(diagnosis_data, args.output)
    else:
        print("Usage:")
        print("  Generate from data:  python generate-report.py --data diagnosis.json --output report.xlsx")
        print("  Generate template:   python generate-report.py --template --output template.xlsx")
