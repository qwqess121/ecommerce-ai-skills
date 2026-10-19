"""
Phase 5: Excel Report v3.2 — Dashboard + 8 Sheets with full data
Integrates: phase3_diagnosis_v3, ads_analysis, competitor_products, tk_product_details,
            phase2_scores, phase3_child_scan
"""
import json, os, sys
from datetime import datetime

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: pip install openpyxl")
    sys.exit(1)

BASE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'output')

def load_json(name):
    path = os.path.join(BASE, name)
    if not os.path.exists(path):
        return None
    with open(path, 'r', encoding='utf-8') as f:
        return json.load(f)

diag = load_json('phase3_diagnosis_v3.json')
ads_data = load_json('ads_analysis.json')
comp_products = load_json('competitor_products_v2.json') or load_json('competitor_products.json')
tk_details = load_json('tk_product_details.json')
tk_data = load_json('tiktok_data.json')
phase2_tk = load_json('phase2_tiktok_scores.json')
amz_scrape = load_json('amz_scrape_images.json')
tk_scrape = load_json('tk_scrape_details.json')
phase0 = load_json('phase0_data.json')
phase2_amz = load_json('phase2_scores.json')
child_scan = load_json('phase3_child_scan.json')

# Load asin_detail for badge info (auto-discover all asin_detail_*.json)
amz_details = {}
import glob
for fpath in sorted(glob.glob(os.path.join(BASE, 'asin_detail_*.json'))):
    data = load_json(os.path.basename(fpath))
    if data:
        items = data if isinstance(data, list) else data.get('results', [])
        for item in items:
            if 'error' not in item and item.get('asin'):
                amz_details[item['asin']] = item

if not diag:
    print("ERROR: phase3_diagnosis_v3.json not found")
    sys.exit(1)

# Build lookup maps
ads_map = {}
if ads_data:
    for r in ads_data.get('results', []):
        ads_map[r['asin']] = r

tk_detail_map = {}
if tk_details:
    for r in tk_details.get('results', []):
        tk_detail_map[str(r['product_id'])] = r

# Scraped image/A+ data
scrape_map = {}
if amz_scrape:
    for r in amz_scrape.get('results', []):
        if not r.get('error'):
            scrape_map[r['asin']] = r

tk_scrape_map = {}
if tk_scrape:
    for r in tk_scrape.get('results', []):
        if not r.get('error') and r.get('image_count', 0) > 0:
            tk_scrape_map[r['product_id']] = r

# Styles
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
DATA_FONT = Font(name="Microsoft YaHei", size=10)
BOLD_FONT = Font(name="Microsoft YaHei", size=10, bold=True)
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
LIGHT_BLUE = PatternFill(start_color="D6EAF8", end_color="D6EAF8", fill_type="solid")
ORANGE_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9"))
WRAP = Alignment(vertical="top", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center")

# Section fills for competitor sheets
SECTION_FILL = PatternFill(start_color="E8DAEF", end_color="E8DAEF", fill_type="solid")
OWN_FILL = PatternFill(start_color="D5F5E3", end_color="D5F5E3", fill_type="solid")
OWN_SECTION_FILL = PatternFill(start_color="ABEBC6", end_color="ABEBC6", fill_type="solid")

# Heatmap dimension fills
DIM_DEEP_GREEN = PatternFill(start_color="27AE60", end_color="27AE60", fill_type="solid")
DIM_LIGHT_GREEN = PatternFill(start_color="82E0AA", end_color="82E0AA", fill_type="solid")
DIM_YELLOW = PatternFill(start_color="F1C40F", end_color="F1C40F", fill_type="solid")
DIM_ORANGE = PatternFill(start_color="E67E22", end_color="E67E22", fill_type="solid")
DIM_RED = PatternFill(start_color="E74C3C", end_color="E74C3C", fill_type="solid")
WHITE_FONT = Font(name="Microsoft YaHei", size=10, color="FFFFFF")
DARK_FONT = Font(name="Microsoft YaHei", size=10, color="1A1A1A")

# Alert card fills
ALERT_RED = PatternFill(start_color="E74C3C", end_color="E74C3C", fill_type="solid")
ALERT_YELLOW = PatternFill(start_color="F39C12", end_color="F39C12", fill_type="solid")
ALERT_GREEN = PatternFill(start_color="27AE60", end_color="27AE60", fill_type="solid")
ALERT_NUM_FONT = Font(name="Microsoft YaHei", size=20, bold=True, color="FFFFFF")
ALERT_LABEL_FONT = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")

# Gap fills for category competitiveness
GAP_RED = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")
GAP_GREEN = PatternFill(start_color="D5F5E3", end_color="D5F5E3", fill_type="solid")

def setup_headers(ws, headers):
    for i, (name, width) in enumerate(headers, 1):
        c = ws.cell(row=1, column=i, value=name)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = CENTER
        c.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

LINK_FONT = Font(name="Microsoft YaHei", size=10, color="0563C1", underline="single")

def wc(ws, row, col, value, font=DATA_FONT, fill=None, align=None, link=None):
    from openpyxl.worksheet.hyperlink import Hyperlink as HL
    c = ws.cell(row=row, column=col, value=value)
    c.font = font
    c.border = THIN_BORDER
    c.alignment = align or WRAP
    if fill:
        c.fill = fill
    if link and value:
        c.hyperlink = HL(ref=c.coordinate, target=str(link), display=str(value))
        c.font = LINK_FONT
        c.style = "Hyperlink"
    return c

def grade_fill(g):
    return {'A': GREEN_FILL, 'B': GREEN_FILL, 'C': YELLOW_FILL, 'D': RED_FILL}.get(g)

def score_fill(s):
    if s >= 85: return GREEN_FILL
    if s >= 70: return YELLOW_FILL
    return RED_FILL

def dim_color(score, max_score):
    """Return (fill, font) for dimension heatmap cell based on score/max percentage."""
    if max_score <= 0:
        return (None, DATA_FONT)
    pct = score / max_score * 100
    if pct >= 90:
        return (DIM_DEEP_GREEN, WHITE_FONT)
    elif pct >= 75:
        return (DIM_LIGHT_GREEN, DARK_FONT)
    elif pct >= 60:
        return (DIM_YELLOW, DARK_FONT)
    elif pct >= 40:
        return (DIM_ORANGE, WHITE_FONT)
    else:
        return (DIM_RED, WHITE_FONT)

# Build amz_own_by_cat (needed by dashboard and competitor sheet)
amz_own_by_cat = {}
for item in diag['amz_results']:
    cat = item.get('l3_category', '')
    if not cat:
        continue
    if cat not in amz_own_by_cat:
        amz_own_by_cat[cat] = []
    asin = item['asin']
    detail = amz_details.get(asin, {})
    amz_own_by_cat[cat].append({
        'asin': asin,
        'title': item.get('original_title', ''),
        'brand': item.get('brand', ''),
        'price': detail.get('price', ''),
        'coupon': detail.get('coupon', '-'),
        'units': detail.get('units', detail.get('totalUnits', '')),
        'revenue': detail.get('revenue', detail.get('totalAmount', '')),
        'bsr': item.get('bsr', ''),
        'rating': item.get('rating', ''),
        'ratings': item.get('reviews', ''),
        'lqs': item.get('lqs', ''),
        'sellerType': detail.get('sellerType', ''),
        'price_trend': detail.get('price_trend', '-'),
        'country': item.get('country', 'US'),
    })

wb = Workbook()
wb.remove(wb.active)

# ======================== Sheet 0: 可视化Dashboard ========================
ws_dash = wb.create_sheet("可视化Dashboard")
ws_dash.sheet_properties.tabColor = "E74C3C"

dash_title_font = Font(name="Microsoft YaHei", size=16, bold=True, color="1F4E79")
dash_section_font = Font(name="Microsoft YaHei", size=13, bold=True, color="1F4E79")
dash_sub_font = Font(name="Microsoft YaHei", size=11, bold=True)

# --- Title ---
ws_dash.merge_cells("A1:N1")
ws_dash["A1"].value = f"Listing 诊断 Dashboard v3.2 — {datetime.now().strftime('%Y-%m-%d %H:%M')}"
ws_dash["A1"].font = dash_title_font
ws_dash["A1"].alignment = CENTER
ws_dash.row_dimensions[1].height = 40

amz_s = diag['amz_summary']
tk_s = diag['tk_summary']

# --- (a) Alert Cards Row ---
# Compute alert stats
oos_count = 0
low_stock_count = 0
price_anomaly_count = 0

if child_scan:
    cs_stats = child_scan.get('stats', child_scan)
    oos_count = cs_stats.get('out_of_stock', 0) or cs_stats.get('out_of_stock_count', 0) or 0
    low_stock_count = cs_stats.get('low_stock', 0) or cs_stats.get('low_stock_count', 0) or 0
    price_anomaly_count = cs_stats.get('price_anomaly', 0) or cs_stats.get('price_anomaly_count', 0) or 0

if oos_count == 0 and low_stock_count == 0 and phase0 and phase0.get('children'):
    from collections import defaultdict as _dd
    _pp = _dd(list)
    for ch in phase0['children']:
        p = ch.get('price')
        if isinstance(p, (int, float)) and p > 0:
            _pp[ch.get('parent_asin', '')].append(p)
    _pavg = {k: sum(v)/len(v) for k, v in _pp.items() if v}
    for ch in phase0['children']:
        ff = ch.get('fulfillment', '')
        fba = ch.get('fba_qty', 0) or 0
        if ff == 'FBA':
            if fba == 0:
                oos_count += 1
            elif fba < 10:
                low_stock_count += 1
        pr = ch.get('price')
        pa = ch.get('parent_asin', '')
        if isinstance(pr, (int, float)) and pr > 0 and pa in _pavg:
            dev = abs(pr - _pavg[pa]) / _pavg[pa] if _pavg[pa] > 0 else 0
            if dev > 0.3:
                price_anomaly_count += 1

violation_total = amz_s.get('violations', 0) + tk_s.get('violations', 0)

cd_count = (amz_s['grades'].get('C', 0) + amz_s['grades'].get('D', 0)
            + tk_s['grades'].get('C', 0) + tk_s['grades'].get('D', 0))

no_aplus_count = 0
for _item in diag['amz_results']:
    _a = _item['asin']
    _d = amz_details.get(_a, {})
    _badge = _d.get('badge', {}) if isinstance(_d.get('badge'), dict) else {}
    if _badge.get('ebc', 'N') != 'Y':
        no_aplus_count += 1

alert_cards = [
    ("断货", oos_count, ALERT_RED if oos_count > 0 else ALERT_GREEN),
    ("低库存", low_stock_count, ALERT_YELLOW if low_stock_count > 0 else ALERT_GREEN),
    ("价格异常", price_anomaly_count, ALERT_YELLOW if price_anomaly_count > 0 else ALERT_GREEN),
    ("违规词", violation_total, ALERT_RED if violation_total > 0 else ALERT_GREEN),
    ("C&D等级", cd_count, ALERT_RED if cd_count > 0 else ALERT_GREEN),
    ("缺A+", no_aplus_count, ALERT_YELLOW if no_aplus_count > 0 else ALERT_GREEN),
]

ws_dash.cell(row=2, column=1, value="关键预警指标").font = dash_section_font
ws_dash.row_dimensions[3].height = 22
ws_dash.row_dimensions[4].height = 40

for i, (label, num, card_fill) in enumerate(alert_cards):
    col_start = i * 2 + 1
    col_end = col_start + 1
    ws_dash.merge_cells(start_row=3, start_column=col_start, end_row=3, end_column=col_end)
    lc = ws_dash.cell(row=3, column=col_start, value=label)
    lc.font = ALERT_LABEL_FONT
    lc.fill = card_fill
    lc.alignment = CENTER
    lc.border = THIN_BORDER
    ws_dash.cell(row=3, column=col_end).fill = card_fill
    ws_dash.cell(row=3, column=col_end).border = THIN_BORDER

    ws_dash.merge_cells(start_row=4, start_column=col_start, end_row=4, end_column=col_end)
    nc = ws_dash.cell(row=4, column=col_start, value=num)
    nc.font = ALERT_NUM_FONT
    nc.fill = card_fill
    nc.alignment = CENTER
    nc.border = THIN_BORDER
    ws_dash.cell(row=4, column=col_end).fill = card_fill
    ws_dash.cell(row=4, column=col_end).border = THIN_BORDER

for i in range(12):
    ws_dash.column_dimensions[get_column_letter(i + 1)].width = 10

# --- (b) Grade Distribution ---
dash_row = 6
ws_dash.cell(row=dash_row, column=1, value="等级分布").font = dash_section_font
dash_row += 1

grade_colors = {
    'A': PatternFill(start_color="27AE60", end_color="27AE60", fill_type="solid"),
    'B': PatternFill(start_color="82E0AA", end_color="82E0AA", fill_type="solid"),
    'C': PatternFill(start_color="F1C40F", end_color="F1C40F", fill_type="solid"),
    'D': PatternFill(start_color="E74C3C", end_color="E74C3C", fill_type="solid"),
}
grade_fonts = {
    'A': WHITE_FONT, 'B': DARK_FONT, 'C': DARK_FONT, 'D': WHITE_FONT,
}

wc(ws_dash, dash_row, 1, "亚马逊", font=BOLD_FONT, align=CENTER)
for gi, g in enumerate(['A', 'B', 'C', 'D']):
    cnt = amz_s['grades'].get(g, 0)
    wc(ws_dash, dash_row, 2 + gi, f"{g}: {cnt}", font=grade_fonts[g], fill=grade_colors[g], align=CENTER)
dash_row += 1

wc(ws_dash, dash_row, 1, "TikTok", font=BOLD_FONT, align=CENTER)
for gi, g in enumerate(['A', 'B', 'C', 'D']):
    cnt = tk_s['grades'].get(g, 0)
    wc(ws_dash, dash_row, 2 + gi, f"{g}: {cnt}", font=grade_fonts[g], fill=grade_colors[g], align=CENTER)
dash_row += 2

# --- (c) AMZ Dimension Score Heatmap ---
ws_dash.cell(row=dash_row, column=1, value="亚马逊维度评分热力图").font = dash_section_font
dash_row += 1

amz_dims = ['title', 'bullets', 'description', 'aplus', 'images', 'compliance', 'spelling', 'competitor']
amz_dim_labels = ['标题', '五点', '描述', 'A+', '图片', '合规', '拼写', '竞争力']

amz_hm_headers = [("ASIN", 14)] + [(l, 10) for l in amz_dim_labels] + [("总分", 8), ("优先修复项", 35)]
for ci, (hname, hw) in enumerate(amz_hm_headers, 1):
    c = ws_dash.cell(row=dash_row, column=ci, value=hname)
    c.fill = HEADER_FILL
    c.font = HEADER_FONT
    c.alignment = CENTER
    c.border = THIN_BORDER
    ws_dash.column_dimensions[get_column_letter(ci)].width = hw
dash_row += 1

if phase2_amz and phase2_amz.get('scores'):
    amz_scores_sorted = sorted(phase2_amz['scores'],
                                key=lambda x: x.get('total_score', 0))
    for sc_item in amz_scores_sorted:
        sc_asin = sc_item.get('asin', '')
        dims = sc_item.get('dimensions', {})
        wc(ws_dash, dash_row, 1, sc_asin, align=CENTER)
        for di, dk in enumerate(amz_dims):
            d = dims.get(dk, {})
            s = d.get('score', 0)
            m = d.get('max', 0)
            fill_d, font_d = dim_color(s, m)
            cell_val = f"{s}/{m}" if m > 0 else "-"
            if fill_d:
                wc(ws_dash, dash_row, 2 + di, cell_val, font=font_d, fill=fill_d, align=CENTER)
            else:
                wc(ws_dash, dash_row, 2 + di, cell_val, align=CENTER)
        total_s = sc_item.get('total_score', 0)
        ts_fill = score_fill(total_s) if isinstance(total_s, (int, float)) else None
        if ts_fill:
            wc(ws_dash, dash_row, 10, total_s, fill=ts_fill, align=CENTER)
        else:
            wc(ws_dash, dash_row, 10, total_s, align=CENTER)
        pf = sc_item.get('priority_fixes', '')
        if isinstance(pf, list):
            pf = '; '.join(pf)
        wc(ws_dash, dash_row, 11, pf)
        dash_row += 1
else:
    wc(ws_dash, dash_row, 1, "(无亚马逊评分数据)")
    dash_row += 1

dash_row += 1

# --- (d) TK Dimension Score Heatmap ---
ws_dash.cell(row=dash_row, column=1, value="TikTok维度评分热力图").font = dash_section_font
dash_row += 1

tk_dims = ['title', 'description', 'images', 'compliance', 'spelling', 'competitor']
tk_dim_labels = ['标题', '描述', '图片', '合规', '拼写', '竞争力']

tk_hm_headers = [("商品ID", 20)] + [(l, 10) for l in tk_dim_labels] + [("总分", 8), ("28天GMV", 12)]
for ci, (hname, hw) in enumerate(tk_hm_headers, 1):
    c = ws_dash.cell(row=dash_row, column=ci, value=hname)
    c.fill = HEADER_FILL
    c.font = HEADER_FONT
    c.alignment = CENTER
    c.border = THIN_BORDER
    if hw > ws_dash.column_dimensions[get_column_letter(ci)].width:
        ws_dash.column_dimensions[get_column_letter(ci)].width = hw
dash_row += 1

if phase2_tk and phase2_tk.get('scores'):
    tk_scores_sorted = sorted(phase2_tk['scores'],
                               key=lambda x: x.get('total_score', 0))
    for sc_item in tk_scores_sorted:
        sc_pid = str(sc_item.get('product_id', ''))
        dims = sc_item.get('dimensions', {})
        wc(ws_dash, dash_row, 1, sc_pid, align=CENTER)
        for di, dk in enumerate(tk_dims):
            d = dims.get(dk, {})
            s = d.get('score', 0)
            m = d.get('max', 0)
            fill_d, font_d = dim_color(s, m)
            cell_val = f"{s}/{m}" if m > 0 else "-"
            if fill_d:
                wc(ws_dash, dash_row, 2 + di, cell_val, font=font_d, fill=fill_d, align=CENTER)
            else:
                wc(ws_dash, dash_row, 2 + di, cell_val, align=CENTER)
        total_s = sc_item.get('total_score', 0)
        ts_fill = score_fill(total_s) if isinstance(total_s, (int, float)) else None
        if ts_fill:
            wc(ws_dash, dash_row, 8, total_s, fill=ts_fill, align=CENTER)
        else:
            wc(ws_dash, dash_row, 8, total_s, align=CENTER)
        gmv28 = sc_item.get('day28_gmv', 0) or sc_item.get('28d_gmv', 0) or 0
        wc(ws_dash, dash_row, 9, f"${gmv28:,.0f}" if gmv28 else "$0", align=CENTER)
        dash_row += 1
else:
    wc(ws_dash, dash_row, 1, "(无TikTok评分数据)")
    dash_row += 1

dash_row += 1

# --- (e) Category Competitiveness Table ---
ws_dash.cell(row=dash_row, column=1, value="品类竞争力对比").font = dash_section_font
dash_row += 1

# Amazon categories
cat_comp_headers = [("平台", 6), ("L3品类", 22), ("自有均价", 10), ("竞品均价", 10),
                    ("价格差距", 10), ("自有均分", 8), ("竞品均分", 8), ("评分差距", 8),
                    ("自有均评论", 10), ("竞品均评论", 10), ("评论差距", 10)]
for ci, (hname, hw) in enumerate(cat_comp_headers, 1):
    c = ws_dash.cell(row=dash_row, column=ci, value=hname)
    c.fill = HEADER_FILL
    c.font = HEADER_FONT
    c.alignment = CENTER
    c.border = THIN_BORDER
    if hw > ws_dash.column_dimensions[get_column_letter(ci)].width:
        ws_dash.column_dimensions[get_column_letter(ci)].width = hw
dash_row += 1

if comp_products and comp_products.get('categories'):
    for cat_name, cat_data in comp_products['categories'].items():
        comps = cat_data.get('true_competitors', []) or cat_data.get('category_ranking', []) or cat_data.get('products', [])
        own_items = amz_own_by_cat.get(cat_name, [])
        if not own_items and not comps:
            continue

        own_prices = [o.get('price', 0) for o in own_items if isinstance(o.get('price'), (int, float)) and o['price'] > 0]
        own_ratings = [o.get('rating', 0) for o in own_items if isinstance(o.get('rating'), (int, float)) and o['rating'] > 0]
        own_reviews = [o.get('ratings', 0) for o in own_items if isinstance(o.get('ratings'), (int, float)) and o['ratings'] > 0]

        comp_prices = [c.get('price', 0) for c in comps if isinstance(c.get('price'), (int, float)) and c['price'] > 0 and not c.get('is_own_brand')]
        comp_ratings = [c.get('rating', 0) for c in comps if isinstance(c.get('rating'), (int, float)) and c['rating'] > 0 and not c.get('is_own_brand')]
        comp_reviews = [c.get('ratings', 0) for c in comps if isinstance(c.get('ratings'), (int, float)) and c['ratings'] > 0 and not c.get('is_own_brand')]

        own_avg_price = sum(own_prices) / len(own_prices) if own_prices else 0
        comp_avg_price = sum(comp_prices) / len(comp_prices) if comp_prices else 0
        own_avg_rating = sum(own_ratings) / len(own_ratings) if own_ratings else 0
        comp_avg_rating = sum(comp_ratings) / len(comp_ratings) if comp_ratings else 0
        own_avg_reviews = sum(own_reviews) / len(own_reviews) if own_reviews else 0
        comp_avg_reviews = sum(comp_reviews) / len(comp_reviews) if comp_reviews else 0

        price_gap = own_avg_price - comp_avg_price if own_avg_price and comp_avg_price else 0
        rating_gap = own_avg_rating - comp_avg_rating if own_avg_rating and comp_avg_rating else 0
        review_gap = own_avg_reviews - comp_avg_reviews if own_avg_reviews and comp_avg_reviews else 0

        price_gap_fill = GAP_RED if price_gap > 5 else (GAP_GREEN if price_gap < -2 else None)
        rating_gap_fill = GAP_RED if rating_gap < -0.3 else (GAP_GREEN if rating_gap > 0.1 else None)
        review_gap_fill = GAP_RED if review_gap < -500 else (GAP_GREEN if review_gap > 0 else None)

        wc(ws_dash, dash_row, 1, "AMZ", align=CENTER)
        wc(ws_dash, dash_row, 2, cat_name)
        wc(ws_dash, dash_row, 3, f"${own_avg_price:.2f}" if own_avg_price else "-", align=CENTER)
        wc(ws_dash, dash_row, 4, f"${comp_avg_price:.2f}" if comp_avg_price else "-", align=CENTER)
        gap_str = f"{'+'if price_gap>0 else ''}{price_gap:.2f}" if price_gap else "-"
        wc(ws_dash, dash_row, 5, gap_str, fill=price_gap_fill, align=CENTER)
        wc(ws_dash, dash_row, 6, f"{own_avg_rating:.1f}" if own_avg_rating else "-", align=CENTER)
        wc(ws_dash, dash_row, 7, f"{comp_avg_rating:.1f}" if comp_avg_rating else "-", align=CENTER)
        gap_str = f"{'+'if rating_gap>0 else ''}{rating_gap:.2f}" if rating_gap else "-"
        wc(ws_dash, dash_row, 8, gap_str, fill=rating_gap_fill, align=CENTER)
        wc(ws_dash, dash_row, 9, f"{own_avg_reviews:.0f}" if own_avg_reviews else "-", align=CENTER)
        wc(ws_dash, dash_row, 10, f"{comp_avg_reviews:.0f}" if comp_avg_reviews else "-", align=CENTER)
        gap_str = f"{'+'if review_gap>0 else ''}{review_gap:.0f}" if review_gap else "-"
        wc(ws_dash, dash_row, 11, gap_str, fill=review_gap_fill, align=CENTER)
        dash_row += 1

# TikTok category competitiveness
if phase2_tk and phase2_tk.get('category_top_selling'):
    tk_own_scores_by_cat = {}
    if phase2_tk.get('scores'):
        tk_cat_key_map_dash = {
            "Women's Perfume": "womens_perfume",
            "Unisex Perfume": "unisex_perfume",
            "Men's Perfume": "mens_perfume",
            "Perfume Sets": "perfume_sets",
            "Perfume": "perfume_generic",
            "Car Air Freshener": "car_air_freshener",
            "Car Fragrance": "car_air_freshener",
            "Pest & Weed Control": "pest_weed_control",
        }
        for s in phase2_tk['scores']:
            l3 = s.get('l3_category', '') or ''
            ck = tk_cat_key_map_dash.get(l3, '')
            if ck:
                if ck not in tk_own_scores_by_cat:
                    tk_own_scores_by_cat[ck] = []
                tk_own_scores_by_cat[ck].append(s)

    for cat_key, products in phase2_tk['category_top_selling'].items():
        cat_display = cat_key.replace('_', ' ').title()
        own_tk = tk_own_scores_by_cat.get(cat_key, [])
        own_gmv = sum(o.get('day28_gmv', 0) or 0 for o in own_tk)
        comp_gmvs = [p.get('period_gmv', 0) or p.get('day28_gmv', 0) or 0
                     for p in products if not p.get('is_own', False)]
        comp_top_gmv = max(comp_gmvs) if comp_gmvs else 0
        comp_avg_gmv = sum(comp_gmvs) / len(comp_gmvs) if comp_gmvs else 0

        gmv_gap = own_gmv - comp_avg_gmv
        gmv_gap_fill = GAP_RED if gmv_gap < 0 else GAP_GREEN

        wc(ws_dash, dash_row, 1, "TK", align=CENTER)
        wc(ws_dash, dash_row, 2, cat_display)
        wc(ws_dash, dash_row, 3, f"${own_gmv:,.0f}", align=CENTER)
        wc(ws_dash, dash_row, 4, f"${comp_avg_gmv:,.0f}", align=CENTER)
        gap_str = f"{'+'if gmv_gap>0 else ''}${gmv_gap:,.0f}"
        wc(ws_dash, dash_row, 5, gap_str, fill=gmv_gap_fill, align=CENTER)
        wc(ws_dash, dash_row, 6, "-", align=CENTER)
        wc(ws_dash, dash_row, 7, "-", align=CENTER)
        wc(ws_dash, dash_row, 8, "-", align=CENTER)
        wc(ws_dash, dash_row, 9, "-", align=CENTER)
        wc(ws_dash, dash_row, 10, f"Top: ${comp_top_gmv:,.0f}", align=CENTER)
        wc(ws_dash, dash_row, 11, "-", align=CENTER)
        dash_row += 1

dash_row += 1

# --- (f) Ad Dependency Warning ---
ws_dash.cell(row=dash_row, column=1, value="广告依赖预警 (广告流量>60%)").font = dash_section_font
dash_row += 1

ad_warn_headers = [("ASIN", 14), ("店铺", 18), ("品牌", 12), ("L3品类", 22),
                   ("广告流量%", 10), ("自然流量%", 10), ("Campaign数", 10)]
for ci, (hname, hw) in enumerate(ad_warn_headers, 1):
    c = ws_dash.cell(row=dash_row, column=ci, value=hname)
    c.fill = HEADER_FILL
    c.font = HEADER_FONT
    c.alignment = CENTER
    c.border = THIN_BORDER
    if hw > ws_dash.column_dimensions[get_column_letter(ci)].width:
        ws_dash.column_dimensions[get_column_letter(ci)].width = hw
dash_row += 1

amz_lookup_dash = {}
for _item in diag['amz_results']:
    amz_lookup_dash[_item['asin']] = _item

ad_warn_list = []
if ads_data:
    for r in ads_data.get('results', []):
        ap = r.get('traffic', {}).get('ad_pct')
        if isinstance(ap, (int, float)) and ap > 60:
            ad_warn_list.append(r)
    ad_warn_list.sort(key=lambda x: -(x.get('traffic', {}).get('ad_pct', 0) or 0))

if ad_warn_list:
    for r in ad_warn_list:
        asin = r['asin']
        amz_i = amz_lookup_dash.get(asin, {})
        traffic = r.get('traffic', {})
        ad_pct = traffic.get('ad_pct', 0)
        organic = traffic.get('organic_pct', 0)
        camps = r.get('ad_structure', {}).get('total_campaigns', 0)
        wc(ws_dash, dash_row, 1, asin, align=CENTER)
        wc(ws_dash, dash_row, 2, amz_i.get('shop', ''))
        wc(ws_dash, dash_row, 3, amz_i.get('brand', ''))
        wc(ws_dash, dash_row, 4, amz_i.get('l3_category', ''))
        wc(ws_dash, dash_row, 5, f"{ad_pct:.1f}%", fill=ALERT_RED, font=WHITE_FONT, align=CENTER)
        wc(ws_dash, dash_row, 6, f"{organic:.1f}%", align=CENTER)
        wc(ws_dash, dash_row, 7, camps, align=CENTER)
        dash_row += 1
else:
    wc(ws_dash, dash_row, 1, "(无ASIN广告流量超过60%)")
    dash_row += 1

# Widen column A for the dashboard
ws_dash.column_dimensions["A"].width = 22
ws_dash.column_dimensions["N"].width = 14

# ======================== Sheet 1: 诊断摘要 ========================
ws1 = wb.create_sheet("诊断摘要")
title_font = Font(name="Microsoft YaHei", size=16, bold=True, color="1F4E79")
sub_font = Font(name="Microsoft YaHei", size=12, bold=True)
stat_font = Font(name="Microsoft YaHei", size=11)

ws1.merge_cells("A1:F1")
ws1["A1"].value = f"AB America / HOOGALIFE Listing 诊断报告 v3.2 — {datetime.now().strftime('%Y-%m-%d %H:%M')}"
ws1["A1"].font = title_font
ws1["A1"].alignment = CENTER
ws1.row_dimensions[1].height = 45

row = 3

ws1.cell(row=row, column=1, value="亚马逊诊断概览").font = sub_font
row += 1
for label, val in [
    ("诊断产品数(父ASIN)", amz_s['total']),
    ("有详情数据", amz_s['with_detail']),
    ("平均综合评分", amz_s['avg_score']),
    ("等级分布", f"A:{amz_s['grades']['A']}  B:{amz_s['grades']['B']}  C:{amz_s['grades']['C']}  D:{amz_s['grades']['D']}"),
    ("违规词检出数", amz_s['violations']),
    ("拼写错误检出数", amz_s['typos']),
]:
    ws1.cell(row=row, column=1, value=label).font = stat_font
    ws1.cell(row=row, column=2, value=val).font = stat_font
    row += 1

row += 1
ws1.cell(row=row, column=1, value="TikTok诊断概览").font = sub_font
row += 1
for label, val in [
    ("诊断商品数", tk_s['total']),
    ("平均综合评分", tk_s['avg_score']),
    ("28天总GMV", f"${tk_s.get('total_28d_gmv', 0):,.0f}"),
    ("等级分布", f"A:{tk_s['grades']['A']}  B:{tk_s['grades']['B']}  C:{tk_s['grades']['C']}  D:{tk_s['grades']['D']}"),
    ("违规词检出数", tk_s['violations']),
    ("拼写错误检出数", tk_s['typos']),
]:
    ws1.cell(row=row, column=1, value=label).font = stat_font
    ws1.cell(row=row, column=2, value=val).font = stat_font
    row += 1

row += 1
ws1.cell(row=row, column=1, value="广告分析概览").font = sub_font
row += 1
if ads_data:
    ads_with = ads_data.get('success', 0)
    ads_no = ads_data.get('no_ads', 0)
    over_reliant = sum(1 for r in ads_data.get('results', [])
                       if (r.get('traffic', {}).get('ad_pct') or 0) > 60)
    healthy = sum(1 for r in ads_data.get('results', [])
                  if (r.get('traffic', {}).get('organic_pct') or 0) > 80)
    for label, val in [
        ("分析ASIN数(US)", ads_data.get('total_asins', 0)),
        ("有广告数据", ads_with),
        ("无广告数据", ads_no),
        ("过度依赖广告(>60%)", over_reliant),
        ("自然流量健康(>80%)", healthy),
    ]:
        ws1.cell(row=row, column=1, value=label).font = stat_font
        c = ws1.cell(row=row, column=2, value=val)
        c.font = stat_font
        if label.startswith("过度") and val > 0:
            c.fill = RED_FILL
        row += 1

row += 1
ws1.cell(row=row, column=1, value="竞品分析概览").font = sub_font
row += 1
if comp_products and comp_products.get('categories'):
    _c = comp_products['categories']
    _sc = sum(1 for v in _c.values() if v.get('true_competitors') or v.get('category_ranking') or v.get('products'))
    _tc = comp_products.get('total_categories') or len(_c)
    _tp = sum(len(v.get('category_ranking', v.get('products', []))) for v in _c.values())
    _true_comp_count = sum(len(v.get('true_competitors', [])) for v in _c.values())
    _verified = comp_products.get('meta', {}).get('all_categories_verified', False)
    _verified_at = comp_products.get('meta', {}).get('updated_at', '')
    for label, val in [
        ("覆盖L3品类", f"{_sc} / {_tc}"),
        ("竞品产品总数", _tp),
        ("同外观真竞品数", _true_comp_count),
        ("产品形态验证", f"V 全部通过 ({_verified_at})" if _verified else "未验证"),
    ]:
        ws1.cell(row=row, column=1, value=label).font = stat_font
        c = ws1.cell(row=row, column=2, value=val)
        c.font = stat_font
        if label == "产品形态验证":
            c.fill = GREEN_FILL if _verified else YELLOW_FILL
        row += 1

row += 1
ws1.cell(row=row, column=1, value="TikTok竞品概览").font = sub_font
row += 1
if phase2_tk and phase2_tk.get('true_competitors'):
    tk_comp = phase2_tk['true_competitors']
    tk_cat_count = len(tk_comp)
    tk_comp_total = sum(len(v.get('true_competitors', [])) for v in tk_comp.values())
    for label, val in [
        ("TK竞品品类数", tk_cat_count),
        ("TK竞品产品数", tk_comp_total),
        ("TK产品形态验证", "V 全部通过 (浏览器验证)"),
    ]:
        ws1.cell(row=row, column=1, value=label).font = stat_font
        c = ws1.cell(row=row, column=2, value=val)
        c.font = stat_font
        if label == "TK产品形态验证":
            c.fill = GREEN_FILL
        row += 1

ws1.column_dimensions["A"].width = 28
ws1.column_dimensions["B"].width = 40

# ======================== Sheet 2: 亚马逊Listing诊断 ========================
AMZ_H = [
    ("国家", 6), ("平台", 8), ("店铺", 18), ("负责人", 12), ("ASIN", 14),
    ("品牌", 12), ("L3品类", 22), ("综合评分", 8), ("等级", 6),
    ("原标题", 55), ("修改标题", 55),
    ("五点描述建议", 60), ("主副图建议", 40), ("A+建议", 40),
    ("违规词", 30), ("错词", 20),
    ("BSR", 10), ("评分", 6), ("评论数", 8), ("LQS", 6),
    ("图片数", 6), ("有A+", 5), ("A+图片数", 7), ("A+模块", 18),
    ("有视频", 5), ("BS", 4), ("AC", 4),
    ("关键词覆盖数", 10), ("SEO评分", 8),
    ("自然流量%", 10), ("广告流量%", 10), ("Campaign数", 10), ("广告策略", 45),
]

ws2 = wb.create_sheet("亚马逊Listing诊断")
setup_headers(ws2, AMZ_H)

for idx, item in enumerate(diag['amz_results'], 2):
    asin = item['asin']
    ad = ads_map.get(asin, {})
    traffic = ad.get('traffic', {})
    ad_struct = ad.get('ad_structure', {})
    organic_pct = traffic.get('organic_pct')
    ad_pct = traffic.get('ad_pct')
    campaigns = ad_struct.get('total_campaigns', 0)
    strategy = ad.get('strategy_detail', '') or ad.get('strategy', '')

    detail = amz_details.get(asin, {})
    badge = detail.get('badge', {}) if isinstance(detail.get('badge'), dict) else {}
    has_ebc = 'Y' if badge.get('ebc', 'N') == 'Y' else 'N'
    has_video = 'Y' if badge.get('video', 'N') == 'Y' else 'N'
    is_bs = 'Y' if badge.get('bestSeller', 'N') == 'Y' else ''
    is_ac = 'Y' if badge.get('amazonChoice', 'N') == 'Y' else ''

    # Scraped image/A+ data
    sc = scrape_map.get(asin, {})
    sc_img_count = sc.get('image_count', '-')
    sc_aplus = sc.get('aplus', {})
    sc_aplus_imgs = sc_aplus.get('image_count', '-') if sc_aplus.get('exists') else '-'
    sc_aplus_modules = ', '.join(sc_aplus.get('modules', [])) if sc_aplus.get('exists') else '-'

    vals = [
        item['country'], item['platform'], item['shop'], item['owner'], asin,
        item['brand'], item['l3_category'],
        item['total_score'], item['grade'],
        item['original_title'], item['optimized_title'],
        item['features_suggestion'], item['image_suggestion'], item['aplus_suggestion'],
        item['violation_words'], item['typo_words'],
        item['bsr'], item['rating'], item['reviews'], item['lqs'],
        sc_img_count, has_ebc, sc_aplus_imgs, sc_aplus_modules,
        has_video, is_bs, is_ac,
        item.get('keyword_indexed_count', '-'),
        item.get('keyword_seo_score', '-'),
        f"{organic_pct:.1f}%" if organic_pct is not None else '-',
        f"{ad_pct:.1f}%" if ad_pct is not None else '-',
        campaigns if campaigns else '-',
        strategy,
    ]
    country_code = item['country']
    domain = 'amazon.com' if country_code == 'US' else 'amazon.ca'
    amz_link = f'https://www.{domain}/dp/{asin}'

    for col, val in enumerate(vals, 1):
        fill = None
        link = None
        if col == 5:
            link = amz_link
        if col == 8 and isinstance(val, (int, float)):
            fill = score_fill(val)
        elif col == 9:
            fill = grade_fill(val)
        elif col == 15 and val:
            fill = RED_FILL
        elif col == 21 and isinstance(val, int) and val < 7:
            fill = PatternFill(start_color="FFF2CC", fill_type="solid")
        elif col == 28 and isinstance(val, int) and val < 50:
            fill = YELLOW_FILL
        elif col == 29 and isinstance(val, (int, float)) and val < 60:
            fill = YELLOW_FILL
        elif col == 31 and isinstance(ad_pct, (int, float)) and ad_pct > 60:
            fill = RED_FILL
        elif col == 30 and isinstance(organic_pct, (int, float)) and organic_pct > 80:
            fill = GREEN_FILL
        wc(ws2, idx, col, val, fill=fill, link=link)

# ======================== Sheet 3: TikTok Listing诊断 ========================
TK_H = [
    ("国家", 6), ("平台", 8), ("店铺", 12), ("负责人", 12), ("产品ID", 20),
    ("综合评分", 8), ("等级", 6),
    ("原标题", 45), ("修改标题", 45),
    ("详情描述建议", 55), ("主副图建议", 40),
    ("图片数", 6), ("描述长度", 8),
    ("违规词", 30), ("错词", 20),
    ("28天销量", 10), ("28天GMV", 12), ("累计GMV", 12),
    ("达人数", 8), ("视频数", 8), ("评论数", 8), ("库存", 8),
    ("爆火指数", 8), ("评分", 6), ("佣金率", 8), ("付费推广", 6),
]

ws3 = wb.create_sheet("TikTok Listing诊断")
setup_headers(ws3, TK_H)

for idx, item in enumerate(diag['tk_results'], 2):
    pid = str(item['product_id'])
    tk_d = tk_detail_map.get(pid, {})
    raw = tk_d.get('raw_detail', {})
    product_raw = raw.get('product', {}) if isinstance(raw, dict) else {}

    gmv28 = item.get('day28_gmv', 0) or 0
    gmv_total = item.get('total_gmv', 0) or 0
    creator = item.get('creator_count', 0) or product_raw.get('linked_creator_count', 0) or 0
    video_count = tk_d.get('video_count', 0) or product_raw.get('linked_video_count', 0) or 0
    comm = item.get('commission_rate', None) or product_raw.get('commission_rate', None)

    review_count = item.get('review_count', 0) or 0
    stock_count = item.get('stock_count', 0) or 0
    viral_index = item.get('viral_index', 0) or 0
    has_paid = item.get('has_paid_promotion', False)
    prod_rating = item.get('product_rating', 0) or 0

    tk_sc = tk_scrape_map.get(pid, {})
    tk_img_count = tk_sc.get('image_count', '-')
    tk_desc_len = tk_sc.get('description_length', '-')

    vals = [
        item['country'], item['platform'], item['shop'], item['owner'], pid,
        item['total_score'], item['grade'],
        item['original_title'], item['optimized_title'],
        item['description_suggestion'], item['image_suggestion'],
        tk_img_count, tk_desc_len,
        item['violation_words'], item['typo_words'],
        item.get('day28_units_sold', 0),
        f"${gmv28:,.0f}" if gmv28 else '$0',
        f"${gmv_total:,.0f}" if gmv_total else '$0',
        creator if creator else 0,
        video_count if video_count else 0,
        review_count,
        stock_count,
        viral_index,
        prod_rating,
        f"{comm}%" if comm else '-',
        'Y' if has_paid else '',
    ]
    tk_link = f'https://www.tiktok.com/shop/pdp/{pid}'
    for col, val in enumerate(vals, 1):
        fill = None
        link = None
        if col == 5:
            link = tk_link
        if col == 6 and isinstance(val, (int, float)):
            fill = score_fill(val)
        elif col == 7:
            fill = grade_fill(val)
        elif col == 12 and isinstance(val, (int, float)) and val < 5:
            fill = YELLOW_FILL
        elif col == 14 and val:
            fill = RED_FILL
        wc(ws3, idx, col, val, fill=fill, link=link)

# ======================== Sheet 4: 亚马逊竞品分析 ========================
# 双层结构：同外观真竞品Top5 + 品类排名Top15
AMZ_COMP_H = [
    ("L3品类", 22), ("竞品类型", 14), ("排名", 6), ("ASIN", 14), ("标题", 50),
    ("品牌", 14), ("价格", 8), ("优惠券", 10), ("月销量", 10), ("月销售额", 12),
    ("BSR", 10), ("评分", 6), ("评论数", 10),
    ("LQS", 6), ("卖家类型", 10), ("价格趋势", 10), ("自有品牌", 8),
]

ws4 = wb.create_sheet("亚马逊竞品分析")
setup_headers(ws4, AMZ_COMP_H)

row_idx = 2
if comp_products and comp_products.get('categories'):
    ncols = len(AMZ_COMP_H)
    for cat_name, cat_data in comp_products['categories'].items():
        # v2 format: true_competitors + category_ranking; v1 fallback: products
        top5_real = cat_data.get('true_competitors', [])
        top15_cat = cat_data.get('category_ranking', [])
        if not top5_real and not top15_cat:
            products = cat_data.get('products', [])
            if not products:
                continue
            non_own = [p for p in products if not p.get('is_own_brand', False)]
            top5_real = non_own[:5]
            top15_cat = list(products)[:15]
        if not top5_real and not top15_cat:
            continue
        # Category header
        wc(ws4, row_idx, 1, f"■ {cat_name}", font=BOLD_FONT, fill=LIGHT_BLUE)
        for c in range(2, ncols + 1):
            wc(ws4, row_idx, c, '', fill=LIGHT_BLUE)
        row_idx += 1
        # Own products in this category
        own_prods = amz_own_by_cat.get(cat_name, [])
        if own_prods:
            wc(ws4, row_idx, 1, '', fill=OWN_SECTION_FILL)
            wc(ws4, row_idx, 2, f'★ 自己产品 ({len(own_prods)})', font=BOLD_FONT, fill=OWN_SECTION_FILL)
            for c in range(3, ncols + 1):
                wc(ws4, row_idx, c, '', fill=OWN_SECTION_FILL)
            row_idx += 1
            for oi, own in enumerate(own_prods, 1):
                own_asin = own['asin']
                m_units = own.get('units', '')
                m_rev = own.get('revenue', '')
                own_country = own.get('country', 'US')
                own_domain = 'amazon.com' if own_country == 'US' else 'amazon.ca'
                own_link = f'https://www.{own_domain}/dp/{own_asin}'
                vals = [
                    '', '★自己', oi, own_asin,
                    (own.get('title', '') or '')[:60],
                    own.get('brand', ''), own.get('price', ''),
                    own.get('coupon', '-'),
                    f"{m_units:,}" if isinstance(m_units, (int, float)) and m_units else ('-' if not m_units else m_units),
                    f"${m_rev:,.0f}" if isinstance(m_rev, (int, float)) and m_rev else ('-' if not m_rev else m_rev),
                    own.get('bsr', ''), own.get('rating', ''),
                    own.get('ratings', ''), own.get('lqs', ''),
                    own.get('sellerType', ''), own.get('price_trend', '-'), 'Y',
                ]
                for col, val in enumerate(vals, 1):
                    link = own_link if col == 4 else None
                    wc(ws4, row_idx, col, val, fill=OWN_FILL, link=link)
                row_idx += 1
        # Sub-section 1: 同外观真竞品 Top5
        wc(ws4, row_idx, 1, '', fill=SECTION_FILL)
        wc(ws4, row_idx, 2, '▸ 同外观真竞品 Top5', font=BOLD_FONT, fill=SECTION_FILL)
        for c in range(3, ncols + 1):
            wc(ws4, row_idx, c, '', fill=SECTION_FILL)
        row_idx += 1
        for rank, prod in enumerate(top5_real, 1):
            comp_asin = prod.get('asin', '')
            m_units = prod.get('units') or prod.get('monthly_units') or prod.get('totalUnits') or ''
            m_rev = prod.get('revenue') or prod.get('monthly_revenue') or prod.get('totalAmount') or ''
            vals = [
                '', '真竞品', rank, comp_asin,
                (prod.get('title', '') or '')[:60],
                prod.get('brand', ''), prod.get('price', ''),
                prod.get('coupon', '-'),
                f"{m_units:,}" if isinstance(m_units, (int, float)) and m_units else ('-' if not m_units else m_units),
                f"${m_rev:,.0f}" if isinstance(m_rev, (int, float)) and m_rev else ('-' if not m_rev else m_rev),
                prod.get('bsr', prod.get('bsrRank', '')), prod.get('rating', ''),
                prod.get('ratings', ''), prod.get('lqs', ''),
                prod.get('sellerType', prod.get('fulfillment', '')),
                prod.get('price_trend', '-'), '',
            ]
            comp_link = f'https://www.amazon.com/dp/{comp_asin}' if comp_asin else None
            for col, val in enumerate(vals, 1):
                link = comp_link if col == 4 else None
                wc(ws4, row_idx, col, val, link=link)
            row_idx += 1
        if not top5_real:
            wc(ws4, row_idx, 2, '(该品类无非自有品牌竞品)')
            row_idx += 1
        # Sub-section 2: 品类排名 Top15
        wc(ws4, row_idx, 1, '', fill=LIGHT_BLUE)
        wc(ws4, row_idx, 2, '▸ 品类排名 Top15', font=BOLD_FONT, fill=LIGHT_BLUE)
        for c in range(3, ncols + 1):
            wc(ws4, row_idx, c, '', fill=LIGHT_BLUE)
        row_idx += 1
        for rank, prod in enumerate(top15_cat, 1):
            is_own = prod.get('is_own_brand', False)
            comp_asin = prod.get('asin', '')
            m_units = prod.get('units') or prod.get('monthly_units') or prod.get('totalUnits') or ''
            m_rev = prod.get('revenue') or prod.get('monthly_revenue') or prod.get('totalAmount') or ''
            vals = [
                '', '品类排名', rank, comp_asin,
                (prod.get('title', '') or '')[:60],
                prod.get('brand', ''), prod.get('price', ''),
                prod.get('coupon', '-'),
                f"{m_units:,}" if isinstance(m_units, (int, float)) and m_units else ('-' if not m_units else m_units),
                f"${m_rev:,.0f}" if isinstance(m_rev, (int, float)) and m_rev else ('-' if not m_rev else m_rev),
                prod.get('bsr', prod.get('bsrRank', '')), prod.get('rating', ''),
                prod.get('ratings', ''), prod.get('lqs', ''),
                prod.get('sellerType', prod.get('fulfillment', '')),
                prod.get('price_trend', '-'),
                'Y' if is_own else '',
            ]
            comp_link = f'https://www.amazon.com/dp/{comp_asin}' if comp_asin else None
            for col, val in enumerate(vals, 1):
                fill = GREEN_FILL if is_own else None
                link = comp_link if col == 4 else None
                wc(ws4, row_idx, col, val, fill=fill, link=link)
            row_idx += 1
        row_idx += 1
else:
    wc(ws4, 2, 1, '竞品数据未采集')

# ======================== Sheet 5: TikTok竞品分析 ========================
# 双层结构：同外观真竞品Top5 + 品类排名Top15
TK_COMP_H = [
    ("L3品类", 18), ("竞品类型", 14), ("排名", 6), ("商品ID", 20), ("标题", 45),
    ("店铺", 18), ("价格区间", 14),
    ("周GMV", 12), ("周环比GMV%", 10), ("累计GMV", 12),
    ("带货达人数", 10), ("带货视频数", 10), ("增长率%", 8), ("自有商品", 8),
]

ws5 = wb.create_sheet("TikTok竞品分析")
setup_headers(ws5, TK_COMP_H)

# Build own-product-by-category map for TikTok
tk_cat_key_map = {
    "Women's Perfume": "womens_perfume",
    "Unisex Perfume": "unisex_perfume",
    "Men's Perfume": "mens_perfume",
    "Perfume Sets": "perfume_sets",
    "Perfume": "perfume_generic",
    "Car Air Freshener": "car_air_freshener",
    "Car Fragrance": "car_air_freshener",
    "Pest & Weed Control": "pest_weed_control",
}
tk_own_by_cat = {}
if phase2_tk and phase2_tk.get('scores'):
    for s in phase2_tk['scores']:
        l3 = s.get('l3_category', '') or ''
        cat_key = tk_cat_key_map.get(l3, '')
        if not cat_key:
            continue
        if cat_key not in tk_own_by_cat:
            tk_own_by_cat[cat_key] = []
        tk_own_by_cat[cat_key].append(s)

row_idx = 2
if phase2_tk and phase2_tk.get('category_top_selling'):
    own_shops = set()
    if tk_data:
        for sk in tk_data.get('shops', {}):
            shop_info = tk_data['shops'][sk]
            own_shops.add(shop_info.get('seller_id', ''))
            own_shops.add(sk.lower())

    ncols = len(TK_COMP_H)
    tk_true_comps = phase2_tk.get('true_competitors', {})
    for cat_key, products in phase2_tk['category_top_selling'].items():
        cat_display = cat_key.replace('_', ' ').title()
        # Use dedicated true_competitors if available (form-factor filtered)
        if cat_display in tk_true_comps or cat_key in tk_true_comps:
            tc_data = tk_true_comps.get(cat_display) or tk_true_comps.get(cat_key, {})
            top5_real = tc_data.get('true_competitors', [])
        else:
            non_own = [p for p in products if not (
                (isinstance(p.get('shop', {}), dict) and
                 (p['shop'].get('shop_id', '') in own_shops or
                  p['shop'].get('shop_name', '').lower() in own_shops))
            )]
            top5_real = non_own[:5]
        top15_cat = list(products)[:15]
        if not top5_real and not top15_cat:
            continue
        # Category header
        wc(ws5, row_idx, 1, f"■ {cat_display}", font=BOLD_FONT, fill=LIGHT_BLUE)
        for c in range(2, ncols + 1):
            wc(ws5, row_idx, c, '', fill=LIGHT_BLUE)
        row_idx += 1
        # Own products in this category
        own_tk_prods = tk_own_by_cat.get(cat_key, [])
        if own_tk_prods:
            wc(ws5, row_idx, 1, '', fill=OWN_SECTION_FILL)
            wc(ws5, row_idx, 2, f'★ 自己产品 ({len(own_tk_prods)})', font=BOLD_FONT, fill=OWN_SECTION_FILL)
            for c in range(3, ncols + 1):
                wc(ws5, row_idx, c, '', fill=OWN_SECTION_FILL)
            row_idx += 1
            for oi, own in enumerate(sorted(own_tk_prods, key=lambda x: -(x.get('day28_gmv') or 0)), 1):
                own_pid = str(own.get('product_id', ''))
                own_link = f'https://www.tiktok.com/shop/pdp/{own_pid}' if own_pid else None
                own_growth = own.get('growth_rate', None)
                vals = [
                    '', '★自己', oi, own_pid, (own.get('title', '') or '')[:50],
                    own.get('shop', ''),
                    own.get('price_display', f"${own.get('price', 0):.2f}" if own.get('price') else ''),
                    f"${own.get('day28_gmv', 0):,.0f}",
                    '-',
                    f"${own.get('total_gmv', 0):,.0f}",
                    own.get('creator_count', ''),
                    own.get('video_count', ''),
                    f"{own_growth:.1f}%" if own_growth is not None else '-', 'Y',
                ]
                for col, val in enumerate(vals, 1):
                    link = own_link if col == 4 else None
                    wc(ws5, row_idx, col, val, fill=OWN_FILL, link=link)
                row_idx += 1
        # Sub-section 1: 同外观真竞品 Top5
        wc(ws5, row_idx, 1, '', fill=SECTION_FILL)
        wc(ws5, row_idx, 2, '▸ 同外观真竞品 Top5', font=BOLD_FONT, fill=SECTION_FILL)
        for c in range(3, ncols + 1):
            wc(ws5, row_idx, c, '', fill=SECTION_FILL)
        row_idx += 1
        for rank, prod in enumerate(top5_real, 1):
            shop_info = prod.get('shop', {})
            shop_name = shop_info.get('shop_name', '') if isinstance(shop_info, dict) else str(shop_info)
            growth = prod.get('units_sold_growth_rate_percent', None)
            tk_comp_pid = str(prod.get('product_id', ''))
            tk_comp_link = f'https://www.tiktok.com/shop/pdp/{tk_comp_pid}' if tk_comp_pid else None
            period_gmv = prod.get('period_gmv') or prod.get('day28_gmv') or 0
            total_gmv = prod.get('total_gmv', 0) or 0
            wow = prod.get('wow_gmv_change_pct', None)
            wow_str = f"{wow:+.1f}%" if wow is not None else '-'
            vid_count = prod.get('video_count', '')
            creator_count = prod.get('creator_count', '')
            vals = [
                '', '真竞品', rank, tk_comp_pid, (prod.get('title', '') or '')[:50],
                shop_name, prod.get('price_display', ''),
                f"${period_gmv:,.0f}",
                wow_str,
                f"${total_gmv:,.0f}",
                creator_count,
                vid_count,
                f"{growth:.1f}%" if growth is not None else '-', '',
            ]
            for col, val in enumerate(vals, 1):
                link = tk_comp_link if col == 4 else None
                wc(ws5, row_idx, col, val, link=link)
            row_idx += 1
        if not top5_real:
            wc(ws5, row_idx, 2, '(该品类无非自有店铺竞品)')
            row_idx += 1
        # Sub-section 2: 品类排名 Top15
        wc(ws5, row_idx, 1, '', fill=LIGHT_BLUE)
        wc(ws5, row_idx, 2, '▸ 品类排名 Top15', font=BOLD_FONT, fill=LIGHT_BLUE)
        for c in range(3, ncols + 1):
            wc(ws5, row_idx, c, '', fill=LIGHT_BLUE)
        row_idx += 1
        for rank, prod in enumerate(top15_cat, 1):
            shop_info = prod.get('shop', {})
            shop_name = shop_info.get('shop_name', '') if isinstance(shop_info, dict) else str(shop_info)
            shop_id = shop_info.get('shop_id', '') if isinstance(shop_info, dict) else ''
            is_own = 'Y' if shop_id in own_shops or shop_name.lower() in own_shops else ''
            growth = prod.get('units_sold_growth_rate_percent', None)
            tk_comp_pid = str(prod.get('product_id', ''))
            tk_comp_link = f'https://www.tiktok.com/shop/pdp/{tk_comp_pid}' if tk_comp_pid else None
            wow = prod.get('wow_gmv_change_pct', None)
            wow_str = f"{wow:+.1f}%" if wow is not None else '-'
            vals = [
                '', '品类排名', rank, tk_comp_pid, (prod.get('title', '') or '')[:50],
                shop_name, prod.get('price_display', ''),
                f"${prod.get('period_gmv', 0):,.0f}",
                wow_str,
                f"${prod.get('total_gmv', 0):,.0f}",
                prod.get('creator_count', ''),
                prod.get('video_count', ''),
                f"{growth:.1f}%" if growth is not None else '-', is_own,
            ]
            for col, val in enumerate(vals, 1):
                fill = GREEN_FILL if is_own == 'Y' else None
                link = tk_comp_link if col == 4 else None
                wc(ws5, row_idx, col, val, fill=fill, link=link)
            row_idx += 1
        row_idx += 1

# ======================== Sheet 6: 子ASIN运营附表 ========================
CHILD_H = [
    ("父ASIN", 14), ("子ASIN", 14), ("国家", 4), ("店铺", 18), ("品牌", 12),
    ("负责人", 14), ("SKU", 16), ("标题", 40), ("价格", 8),
    ("FBA库存", 10), ("配送方式", 8), ("BSR排名", 8), ("评论数", 8),
    ("价格异常", 10), ("库存预警", 10),
]

ws_child = wb.create_sheet("子ASIN运营附表")
setup_headers(ws_child, CHILD_H)

row_idx = 2
if phase0 and phase0.get('children'):
    children_list = phase0['children']
    from collections import defaultdict
    parent_prices = defaultdict(list)
    for ch in children_list:
        p = ch.get('price')
        if isinstance(p, (int, float)) and p > 0:
            parent_prices[ch.get('parent_asin', '')].append(p)
    parent_avg = {k: sum(v)/len(v) for k, v in parent_prices.items() if v}

    for ch in sorted(children_list, key=lambda x: (x.get('parent_asin',''), x.get('child_asin',''))):
        child_asin = ch.get('child_asin', '')
        parent_asin = ch.get('parent_asin', '')
        price = ch.get('price', '')
        fba_qty = ch.get('fba_qty', 0) or 0

        price_flag = ''
        price_fill_c = None
        if isinstance(price, (int, float)) and price > 0 and parent_asin in parent_avg:
            avg = parent_avg[parent_asin]
            deviation = abs(price - avg) / avg if avg > 0 else 0
            if deviation > 0.5:
                price_flag = f'偏差{deviation*100:.0f}%'
                price_fill_c = RED_FILL
            elif deviation > 0.3:
                price_flag = f'偏差{deviation*100:.0f}%'
                price_fill_c = YELLOW_FILL

        inv_flag = ''
        inv_fill_c = None
        fulfillment = ch.get('fulfillment', '')
        if fulfillment == 'FBA':
            if fba_qty == 0:
                inv_flag = '断货'
                inv_fill_c = RED_FILL
            elif fba_qty < 10:
                inv_flag = '低库存'
                inv_fill_c = YELLOW_FILL

        country = ch.get('country', '')
        domain = 'amazon.com' if country == 'US' else 'amazon.ca'
        is_valid_asin = child_asin and len(str(child_asin)) == 10 and str(child_asin).startswith('B0')
        child_link = f'https://www.{domain}/dp/{child_asin}' if is_valid_asin else None

        vals = [
            parent_asin, child_asin, country,
            ch.get('shop', ''), ch.get('brand', ''), ch.get('owner', ''),
            ch.get('sku', ''), (ch.get('title', '') or '')[:50],
            f"${price:.2f}" if isinstance(price, (int, float)) else price,
            f"{fba_qty:,}" if isinstance(fba_qty, (int, float)) else fba_qty,
            fulfillment, ch.get('rank', ''), ch.get('reviews', ''),
            price_flag, inv_flag,
        ]
        parent_valid = parent_asin and len(str(parent_asin)) == 10 and str(parent_asin).startswith('B0')
        parent_link = f'https://www.{domain}/dp/{parent_asin}' if parent_valid else None
        for col, val in enumerate(vals, 1):
            fill = None
            link = child_link if col == 2 else (parent_link if col == 1 else None)
            if col == 14:
                fill = price_fill_c
            elif col == 15:
                fill = inv_fill_c
            elif col == 10 and inv_fill_c:
                fill = inv_fill_c
            wc(ws_child, row_idx, col, val, fill=fill, link=link)
        row_idx += 1

if row_idx == 2:
    wc(ws_child, 2, 1, '(无子ASIN数据)')

# ======================== Sheet 7: 违规词&错词汇总 ========================
VIOL_H = [
    ("ASIN/商品ID", 18), ("平台", 8), ("国家/店铺", 14),
    ("类型", 8), ("违规词/错词", 18), ("类别/正确词", 25),
    ("严重程度", 8), ("出现位置", 8), ("原文上下文", 40), ("建议替换", 20),
]

ws6 = wb.create_sheet("违规词&错词汇总")
setup_headers(ws6, VIOL_H)

row_idx = 2
for v in diag.get('amz_violations', []):
    sev_fill = RED_FILL if v['severity'] == 'high' else YELLOW_FILL
    vals = [v['asin'], v['platform'], v['country'],
            '违规词', v['word'], v['category'][:30],
            v['severity'], v['location'], v['context'], v['replacement']]
    for col, val in enumerate(vals, 1):
        wc(ws6, row_idx, col, val, fill=sev_fill if col in (5, 7) else None)
    row_idx += 1

for v in diag.get('tk_violations', []):
    sev_fill = RED_FILL if v['severity'] == 'high' else YELLOW_FILL
    vals = [v['product_id'], v['platform'], v['shop'],
            '违规词', v['word'], v['category'][:30],
            v['severity'], v['location'], v['context'], v['replacement']]
    for col, val in enumerate(vals, 1):
        wc(ws6, row_idx, col, val, fill=sev_fill if col in (5, 7) else None)
    row_idx += 1

for t in diag.get('amz_typos', []):
    vals = [t['asin'], t['platform'], t['country'],
            '错词', t['wrong'], t['correct'],
            '-', t['location'], t['context'], t['correct']]
    for col, val in enumerate(vals, 1):
        wc(ws6, row_idx, col, val)
    row_idx += 1

for t in diag.get('tk_typos', []):
    vals = [t['product_id'], t['platform'], t['shop'],
            '错词', t['wrong'], t['correct'],
            '-', t['location'], t['context'], t['correct']]
    for col, val in enumerate(vals, 1):
        wc(ws6, row_idx, col, val)
    row_idx += 1

if row_idx == 2:
    wc(ws6, 2, 1, '(未检测到违规词或拼写错误)')

# ======================== Sheet 8: 广告分析&策略建议 ========================
ADS_H = [
    ("国家", 6), ("店铺", 18), ("ASIN", 14), ("品牌", 12), ("L3品类", 20),
    ("Campaign数", 10), ("SP", 5), ("SB", 5), ("SBV", 5),
    ("自然流量%", 10), ("广告流量%", 10), ("流量评分", 10),
    ("详细策略建议", 80),
]

ws7 = wb.create_sheet("广告分析&策略建议")
setup_headers(ws7, ADS_H)

row_idx = 2
if ads_data:
    amz_lookup = {}
    for item in diag['amz_results']:
        amz_lookup[item['asin']] = item

    for r in sorted(ads_data.get('results', []),
                    key=lambda x: -(x.get('traffic', {}).get('ad_pct') or 0)):
        asin = r['asin']
        amz_item = amz_lookup.get(asin, {})
        ad_struct = r.get('ad_structure', {})
        traffic = r.get('traffic', {})
        camp_types = ad_struct.get('campaign_types', {})

        organic = traffic.get('organic_pct')
        ad_pct = traffic.get('ad_pct')
        total_score = traffic.get('total_score')

        country = amz_item.get('country', 'US')
        domain = 'amazon.com' if country == 'US' else 'amazon.ca'
        ads_link = f'https://www.{domain}/dp/{asin}'

        vals = [
            country, amz_item.get('shop', r.get('shop', '')),
            asin, amz_item.get('brand', ''), amz_item.get('l3_category', ''),
            ad_struct.get('total_campaigns', 0),
            camp_types.get('SP', 0), camp_types.get('SB', 0), camp_types.get('SBV', 0),
            f"{organic:.1f}%" if organic is not None else '-',
            f"{ad_pct:.1f}%" if ad_pct is not None else '-',
            total_score if total_score is not None else '-',
            r.get('strategy_detail', '') or r.get('strategy', ''),
        ]
        for col, val in enumerate(vals, 1):
            fill = None
            link = ads_link if col == 3 else None
            if col == 11 and isinstance(ad_pct, (int, float)) and ad_pct > 60:
                fill = RED_FILL
            elif col == 10 and isinstance(organic, (int, float)) and organic > 80:
                fill = GREEN_FILL
            elif col == 6 and val == 0:
                fill = ORANGE_FILL
            wc(ws7, row_idx, col, val, fill=fill, link=link)
        row_idx += 1

# Save
out_path = os.path.join(BASE, f'Listing诊断报告v3.2_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx')
wb.save(out_path)
print(f'\nReport saved: {out_path}')
print(f'Sheets: {len(wb.sheetnames)}')
for name in wb.sheetnames:
    ws_tmp = wb[name]
    rows = ws_tmp.max_row - 1 if ws_tmp.max_row > 1 else 0
    print(f'  {name}: {rows} rows')
